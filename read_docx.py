import docx

def read_docx(file_path):
    doc = docx.Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells]
            full_text.append(" | ".join(row_text))
    return '\n'.join(full_text)

try:
    content = read_docx("AI_Problem Statement.docx")
    with open("problem_statement.txt", "w", encoding="utf-8") as f:
        f.write(content)
    print("Successfully read docx and wrote to problem_statement.txt")
except Exception as e:
    print(f"Error: {e}")
